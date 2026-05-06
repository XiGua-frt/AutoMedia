from __future__ import annotations

from pathlib import Path
from typing import Iterable

from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import Paragraph

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT_DIR = Path(__file__).resolve().parent
PDF_PATH = OUT_DIR / "吴文博_AI后端_RAG实习_简历.pdf"
DOCX_PATH = OUT_DIR / "吴文博_AI后端_RAG实习_简历.docx"


def _register_fonts() -> None:
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))


def _draw_wrapped(
    canvas,
    text: str,
    x: float,
    y: float,
    width: float,
    font_size: float = 9.4,
    leading: float = 12.5,
    color=colors.HexColor("#252A31"),
    bold: bool = False,
) -> float:
    style = ParagraphStyle(
        "body",
        fontName="STSong-Light",
        fontSize=font_size,
        leading=leading,
        textColor=color,
        alignment=TA_LEFT,
        spaceAfter=0,
        spaceBefore=0,
    )
    if bold:
        text = f"<b>{text}</b>"
    p = Paragraph(text, style)
    _, height = p.wrap(width, 100 * mm)
    p.drawOn(canvas, x, y - height)
    return y - height


def _section(canvas, title: str, y: float, page_width: float) -> float:
    left = 18 * mm
    right = page_width - 18 * mm
    bar_h = 9.5 * mm
    canvas.setFillColor(colors.HexColor("#EEF3FF"))
    canvas.rect(left, y - bar_h, right - left, bar_h, fill=1, stroke=0)
    canvas.setFillColor(colors.HexColor("#3F73F1"))
    canvas.rect(left, y - bar_h, 2.2 * mm, bar_h, fill=1, stroke=0)
    canvas.setFillColor(colors.HexColor("#3F73F1"))
    canvas.setFont("STSong-Light", 16)
    canvas.drawString(left + 7 * mm, y - 6.7 * mm, title)
    return y - bar_h - 6 * mm


def _meta_line(canvas, left: str, right: str, x: float, y: float, w: float) -> None:
    canvas.setFont("STSong-Light", 10)
    canvas.setFillColor(colors.HexColor("#252A31"))
    canvas.drawString(x, y, left)
    canvas.setFillColor(colors.HexColor("#7A7F87"))
    canvas.drawRightString(x + w, y, right)


def _bullet(canvas, text: str, x: float, y: float, w: float) -> float:
    canvas.setFillColor(colors.HexColor("#3F73F1"))
    canvas.circle(x + 1.2 * mm, y - 3.2, 1.1, fill=1, stroke=0)
    return _draw_wrapped(canvas, text, x + 4.2 * mm, y + 1, w - 4.2 * mm, 8.7, 11.2)


def build_pdf() -> None:
    _register_fonts()
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(PDF_PATH), pagesize=A4)
    width, height = A4
    left = 18 * mm
    right = width - 18 * mm
    content_w = right - left
    y = height - 18 * mm

    c.setFillColor(colors.HexColor("#20242A"))
    c.setFont("STSong-Light", 24)
    c.drawString(left, y, "吴文博")
    c.setFont("STSong-Light", 10.5)
    c.setFillColor(colors.HexColor("#3F73F1"))
    c.drawString(left, y - 8.5 * mm, "AI 后端 / RAG / Agent 工程实习")
    c.setFillColor(colors.HexColor("#2D333B"))
    c.setFont("STSong-Light", 9.5)
    c.drawString(left, y - 16 * mm, "电话：待补充  |  邮箱：待补充  |  技术博客：知乎 / 公众号累计阅读 5.5w+")

    c.setStrokeColor(colors.HexColor("#D7E2FF"))
    c.setLineWidth(1)
    c.line(left, y - 22 * mm, right, y - 22 * mm)

    y -= 31 * mm
    y = _section(c, "教育经历", y, width)
    _meta_line(c, "江西理工大学 · 软件工程 · 本科", "2024.09 - 2028.06（预计）", left, y, content_w)
    y -= 6.2 * mm
    _draw_wrapped(c, "GPA 3.3/4.0，专业前 14%；两次三等奖学金；英语六级。", left, y, content_w, 9.3, 12)
    y -= 6.2 * mm
    _draw_wrapped(c, "竞赛：华教杯数学竞赛省级三等奖；蓝桥杯省级三等奖。", left, y, content_w, 9.3, 12)
    y -= 8.5 * mm

    y = _section(c, "专业技能", y, width)
    skills = [
        "<b>语言与工程：</b>熟悉 Python、C++，了解 Java；熟悉 FastAPI 架构、Pydantic、异步服务开发与接口设计。",
        "<b>AI / RAG：</b>熟悉 RAG 入库与检索链路，了解 LangChain、LangGraph、Qdrant、DashScope Embedding、中文 chunk 切分、MMR 与元数据过滤。",
        "<b>后端与基础设施：</b>熟悉 MySQL、Redis、Docker Compose、SSE 实时推送；了解 Vue3 + Vite 前端工程协作。",
        "<b>开发效率：</b>熟练使用 Cursor、Codex 等 AI 编程工具，具备独立拆解需求、阅读文档、实现与联调能力。",
    ]
    for item in skills:
        y = _bullet(c, item, left, y, content_w)
        y -= 2.5 * mm
    y -= 2 * mm

    y = _section(c, "项目经历", y, width)
    _meta_line(c, "AI 爆款文章创作器 · 多智能体 RAG 内容创作平台", "2026.04 - 至今", left, y, content_w)
    y -= 5.8 * mm
    _draw_wrapped(
        c,
        "技术栈：FastAPI、Vue3、MySQL、Redis、Qdrant、LangChain、DashScope、SSE、Docker Compose",
        left,
        y,
        content_w,
        8.8,
        11,
        color=colors.HexColor("#606873"),
    )
    y -= 5.8 * mm
    project_bullets = [
        "围绕“输入主题自动生成图文 Markdown 文章”的场景，设计多智能体创作链路，串联 RAG 检索、标题生成、大纲规划、正文写作、配图分析、并行配图与图文合成等节点。",
        "实现 FastAPI 后端分层架构，配合 MySQL 持久化文章 / 执行日志元数据，Redis 支撑 Session / 缓存，SSE 向前端实时推送 Agent 阶段状态，提升长任务可观测性。",
        "落地 RAG 一期基础设施：封装 DashScope Embedding、Qdrant 向量写入与检索、文档加载、正文清洗、元数据抽取、摘要生成、Markdown Header + 中文递归两级切分。",
        "设计知识库入库 Pipeline，支持 dry-run 联调与真实写入；补充知识文档 / chunk 元数据 SQL 表结构，使向量 payload 与 MySQL 元数据保持一致。",
        "基于现有手写状态共享编排，规划 LangGraph 渐进式升级方案：将 ArticleState 迁移为 TypedDict State，通过 Node 适配函数、条件边、Checkpointer 与 Send API 支持断点续跑、错误路由和并行配图。",
    ]
    for item in project_bullets:
        y = _bullet(c, item, left, y, content_w)
        y -= 2 * mm

    y -= 1 * mm
    _meta_line(c, "MDNote Markdown 笔记编辑工具 · 独立开发", "用户 100+", left, y, content_w)
    y -= 5.5 * mm
    for item in [
        "使用 Cursor 独立完成需求拆解、编辑器功能开发与发布迭代，覆盖 Markdown 编辑、预览与常用笔记工作流。",
        "产品实际使用人数 100+，具备从个人工具到真实用户反馈闭环的完整开发经验。",
    ]:
        y = _bullet(c, item, left, y, content_w)
        y -= 2 * mm

    y -= 2 * mm
    y = _section(c, "技术影响力", y, width)
    for item in [
        "持续输出 AI 工程与 RAG 技术博客，知乎 / 公众号累计阅读量 5.5w+。",
        "RAG chunk 切分、Embedding、检索系列文章阅读量 3.1w+，曾获知乎首页推荐。",
    ]:
        y = _bullet(c, item, left, y, content_w)
        y -= 2.2 * mm

    c.showPage()
    c.save()


def _set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _set_cell_border(cell, color: str = "FFFFFF") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = OxmlElement(f"w:{edge}")
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "0")
        tag.set(qn("w:color"), color)
        borders.append(tag)
    tc_pr.append(borders)


def _docx_run(paragraph, text: str, size: float = 9.5, bold: bool = False, color: str = "252A31"):
    run = paragraph.add_run(text)
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    return run


def _docx_section(doc: Document, title: str) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Cm(0.25)
    table.columns[1].width = Cm(16.7)
    for cell in table.rows[0].cells:
        _set_cell_border(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    _set_cell_shading(table.cell(0, 0), "3F73F1")
    _set_cell_shading(table.cell(0, 1), "EEF3FF")
    p = table.cell(0, 1).paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    _docx_run(p, title, 15, True, "3F73F1")


def _docx_bullets(doc: Document, items: Iterable[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style=None)
        p.paragraph_format.left_indent = Cm(0.2)
        p.paragraph_format.first_line_indent = Cm(-0.2)
        p.paragraph_format.space_after = Pt(2.5)
        p.paragraph_format.line_spacing = 1.05
        _docx_run(p, "• ", 9.2, False, "3F73F1")
        _docx_run(p, item, 9.2)


def build_docx() -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.4)
    section.bottom_margin = Cm(1.2)
    section.left_margin = Cm(1.55)
    section.right_margin = Cm(1.55)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(2)
    _docx_run(title, "吴文博", 24, True, "20242A")
    sub = doc.add_paragraph()
    sub.paragraph_format.space_after = Pt(1)
    _docx_run(sub, "AI 后端 / RAG / Agent 工程实习", 10.5, True, "3F73F1")
    contact = doc.add_paragraph()
    contact.paragraph_format.space_after = Pt(8)
    _docx_run(contact, "电话：待补充  |  邮箱：待补充  |  技术博客：知乎 / 公众号累计阅读 5.5w+", 9.5)

    _docx_section(doc, "教育经历")
    p = doc.add_paragraph()
    _docx_run(p, "江西理工大学 · 软件工程 · 本科", 10.3, True)
    _docx_run(p, "    2024.09 - 2028.06（预计）", 9.5, False, "7A7F87")
    _docx_bullets(doc, ["GPA 3.3/4.0，专业前 14%；两次三等奖学金；英语六级。", "竞赛：华教杯数学竞赛省级三等奖；蓝桥杯省级三等奖。"])

    _docx_section(doc, "专业技能")
    _docx_bullets(
        doc,
        [
            "语言与工程：熟悉 Python、C++，了解 Java；熟悉 FastAPI 架构、Pydantic、异步服务开发与接口设计。",
            "AI / RAG：熟悉 RAG 入库与检索链路，了解 LangChain、LangGraph、Qdrant、DashScope Embedding、中文 chunk 切分、MMR 与元数据过滤。",
            "后端与基础设施：熟悉 MySQL、Redis、Docker Compose、SSE 实时推送；了解 Vue3 + Vite 前端工程协作。",
            "开发效率：熟练使用 Cursor、Codex 等 AI 编程工具，具备独立拆解需求、阅读文档、实现与联调能力。",
        ],
    )

    _docx_section(doc, "项目经历")
    p = doc.add_paragraph()
    _docx_run(p, "AI 爆款文章创作器 · 多智能体 RAG 内容创作平台", 10.3, True)
    _docx_run(p, "    2026.04 - 至今", 9.5, False, "7A7F87")
    p = doc.add_paragraph()
    _docx_run(p, "技术栈：FastAPI、Vue3、MySQL、Redis、Qdrant、LangChain、DashScope、SSE、Docker Compose", 8.7, False, "606873")
    _docx_bullets(
        doc,
        [
            "围绕“输入主题自动生成图文 Markdown 文章”的场景，设计多智能体创作链路，串联 RAG 检索、标题生成、大纲规划、正文写作、配图分析、并行配图与图文合成等节点。",
            "实现 FastAPI 后端分层架构，配合 MySQL 持久化文章 / 执行日志元数据，Redis 支撑 Session / 缓存，SSE 向前端实时推送 Agent 阶段状态，提升长任务可观测性。",
            "落地 RAG 一期基础设施：封装 DashScope Embedding、Qdrant 向量写入与检索、文档加载、正文清洗、元数据抽取、摘要生成、Markdown Header + 中文递归两级切分。",
            "设计知识库入库 Pipeline，支持 dry-run 联调与真实写入；补充知识文档 / chunk 元数据 SQL 表结构，使向量 payload 与 MySQL 元数据保持一致。",
            "基于现有手写状态共享编排，规划 LangGraph 渐进式升级方案：将 ArticleState 迁移为 TypedDict State，通过 Node 适配函数、条件边、Checkpointer 与 Send API 支持断点续跑、错误路由和并行配图。",
        ],
    )

    p = doc.add_paragraph()
    _docx_run(p, "MDNote Markdown 笔记编辑工具 · 独立开发", 10.3, True)
    _docx_run(p, "    用户 100+", 9.5, False, "7A7F87")
    _docx_bullets(
        doc,
        [
            "使用 Cursor 独立完成需求拆解、编辑器功能开发与发布迭代，覆盖 Markdown 编辑、预览与常用笔记工作流。",
            "产品实际使用人数 100+，具备从个人工具到真实用户反馈闭环的完整开发经验。",
        ],
    )

    _docx_section(doc, "技术影响力")
    _docx_bullets(
        doc,
        [
            "持续输出 AI 工程与 RAG 技术博客，知乎 / 公众号累计阅读量 5.5w+。",
            "RAG chunk 切分、Embedding、检索系列文章阅读量 3.1w+，曾获知乎首页推荐。",
        ],
    )

    for paragraph in doc.paragraphs:
        paragraph.paragraph_format.space_before = Pt(0)
    doc.save(DOCX_PATH)


if __name__ == "__main__":
    build_pdf()
    build_docx()
    print(PDF_PATH)
    print(DOCX_PATH)
